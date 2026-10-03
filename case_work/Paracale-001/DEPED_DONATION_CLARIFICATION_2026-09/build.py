# DepEd request packet: 1972 school donation out of TCT T-4185 (Lot 3, Psu-143364, San Rafael, Jose Panganiban).
#   01_DEPED_REQUEST_LETTER_8.5x13.docx/.pdf  -> 14 pt letter, signed by Allan
#   00_DEPED_PACKET_8.5x13.pdf                -> letter + Annex "A" (T-4185 certified true copy, 4 pp)
# Sources: TCT T-4185 CTC, RD Camarines Norte Ref. 2024003960 (doc 494): p1 location/owner, p2 entry PE-39428-102-51.
# Run: python3 build.py   (needs python-docx, PyMuPDF, soffice)
import os, subprocess, fitz
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
TITLE = os.path.join(HERE, "TCT_T-4185_CTC_2024-07-11.pdf")
W, H = 612, 936  # 8.5 x 13 in
FS = 14


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(13)
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.8)
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(FS)
    st.paragraph_format.space_after = Pt(6)
    return d


def para(d, text, bold=False, align=None, size=None, italic=False, indent=None, after=None):
    p = d.add_paragraph()
    for i, part in enumerate(text.split("**")):
        r = p.add_run(part); r.bold = bold or (i % 2 == 1); r.italic = italic
        if size: r.font.size = Pt(size)
    p.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT,
                   "l": WD_ALIGN_PARAGRAPH.LEFT}.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY)
    if indent: p.paragraph_format.left_indent = Inches(indent)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p


def build_letter():
    d = new_doc()
    para(d, "21 September 2026", align="r")
    para(d, "**THE SCHOOLS DIVISION SUPERINTENDENT**\nDepartment of Education\nSchools Division of Camarines Norte\nDaet, Camarines Norte", align="l")
    para(d, "**Re: Request for clarification and records: 10,000 sq m donated in 1972 \"for school purposes only\" "
            "out of TCT No. T-4185, Lot 3, Psu-143364, Barrio San Rafael, Jose Panganiban, Camarines Norte**", align="l")
    para(d, "Dear Sir/Madam:", align="l")
    para(d, "I write as a son and heir of the late spouses Vicente L. Inocalla, Sr. and Beatriz Villafria Inocalla, "
            "for myself and on behalf of my co-heirs.")
    para(d, "**1. The donation.** Transfer Certificate of Title No. T-4185, registered in the name of Beatriz Villafria "
            "and covering Lot 3 of plan Psu-143364 in Barrio San Rafael, Jose Panganiban, carries this entry:")
    para(d, "PE-39428-102-51, DEED OF DONATION, executed in favor of the Bureau of Public Schools, or its authorized "
            "representative in Camarines Norte, **for school purposes only**, of a portion of **10,000 square meters** of the "
            "land described in the title; Doc. No. 127, Page No. 26, Book No. I, Series of 1972, Notary Public "
            "Alberto L. Totones, dated at Jose Panganiban, Camarines Norte, on 1 March 1972; inscribed 12 February 1973.",
         indent=0.5, italic=True)
    para(d, "We understand the Department of Education is the successor of the Bureau of Public Schools as donee.")
    para(d, "**2. Purpose.** The family is updating its property records and wishes to confirm the present status "
            "of the donated portion and how it is being used.")
    para(d, "**3. What we ask.** We respectfully ask your office to state in writing:")
    items = [
        "**(a) Which portion was donated:** its location and boundaries, with any sketch plan, segregation or "
        "subdivision plan, or survey on file; and whether it was ever titled or tax-declared in the name of the "
        "Bureau of Public Schools or DepEd.",
        "**(b) Whether a school was built and still operates:** its name and school ID, the years it operated, and, "
        "if it closed or transferred, when and why.",
        "**(c) Whether the Department has allowed anyone to occupy or use the area:** any permission, lease, "
        "usufruct, agreement or tolerance, for any purpose, and if so to whom, "
        "when and on what terms. If none, please say so.",
        "**(d) Copies of the Deed of Donation** and of any acceptance by the Bureau of Public Schools or DepEd, "
        "with any property inventory record for the site.",
    ]
    for t in items:
        para(d, t, indent=0.4)
    para(d, "**4.** Please send your reply to the address below. We would welcome a site visit by your property "
            "custodian or division engineer, and will point out the area.")
    para(d, "Thank you for your assistance.")
    para(d, "Respectfully,", align="l", after=30)
    para(d, "______________________________\n**ALLAN V. INOCALLA**\nFor himself and on behalf of the heirs of\n"
            "Spouses Vicente L. Inocalla, Sr. and Beatriz Villafria Inocalla\nc/o Golden Inocalla Management Services\nPurok 3, Bagasbas Road, Brgy. Bagasbas, Daet, Camarines Norte\n"
            "Mobile: 0917 155 4782\nE-mail: shiraction2@gmail.com", align="l")
    para(d, "Annex \"A\": Certified true copy of TCT No. T-4185, Registry of Deeds of Camarines Norte, Ref. No. 2024003960, "
            "11 July 2024 (entry PE-39428-102-51 on page 2).", align="l", size=12)
    para(d, "RECEIVED BY: ______________________   Date/Time: ______________", align="l", size=12)
    out = os.path.join(HERE, "01_DEPED_REQUEST_LETTER_8.5x13.docx")
    d.save(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, out], check=True, capture_output=True)
    return out[:-5] + ".pdf"


def build_packet(letter_pdf):
    doc = fitz.open()
    doc.insert_pdf(fitz.open(letter_pdf))
    p = doc.new_page(width=W, height=H)
    p.insert_textbox(fitz.Rect(60, 320, W - 60, 400), 'ANNEX "A"', fontsize=36, fontname="tibo", align=1)
    p.insert_textbox(fitz.Rect(60, 410, W - 60, 480), "Certified True Copy of TCT No. T-4185", fontsize=16, fontname="tibo", align=1)
    p.insert_textbox(fitz.Rect(70, 490, W - 70, 700),
                     "Registry of Deeds of Camarines Norte, Ref. No. 2024003960, 11 July 2024 (4 pages)\n"
                     "Entry PE-39428-102-51, Deed of Donation, is on page 2",
                     fontsize=13, fontname="tiro", align=1)
    s = fitz.open(TITLE)
    for i in range(len(s)):
        p = doc.new_page(width=W, height=H)
        rot = 90 if s[i].rect.width > s[i].rect.height else 0  # landscape source page -> upright
        p.show_pdf_page(fitz.Rect(18, 18, W - 18, H - 18), s, i, keep_proportion=True, rotate=rot)
    out = os.path.join(HERE, "00_DEPED_PACKET_8.5x13.pdf")
    doc.save(out)
    return out


if __name__ == "__main__":
    lp = build_letter()
    print(lp, fitz.open(lp).page_count, "pp")
    pk = build_packet(lp)
    print(pk, fitz.open(pk).page_count, "pp")
