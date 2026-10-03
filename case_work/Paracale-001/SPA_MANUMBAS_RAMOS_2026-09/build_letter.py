# -*- coding: utf-8 -*-
"""Light authorization letter — Allan signs, no notarization. Run: python3 build_letter.py"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER

HEAD = [("b", "ALLAN VILLAFRIA INOCALLA"),
        ("n", "Proprietor · AVI GOLD PROCESSING PLANT (DTI Business Name No. 8480852)"),
        ("n", "Purok 4, Barangay Capacuan, Paracale, Camarines Norte 4605"),
        ("n", "Mobile 0917 155 4782 · shiraction2@gmail.com")]

BODY = [
 "17 September 2026", "",
 "<b>TO WHOM IT MAY CONCERN</b>",
 "All government offices, agencies and local government units concerned", "",
 "<b>SUBJECT: AUTHORIZATION OF REPRESENTATIVES — DOCUMENTS AND PERMITS</b>", "",
  ("I, <b>ALLAN VILLAFRIA INOCALLA</b>, Filipino, of legal age, of the address above, hereby authorize —"),
 "<b>MR. ROMEO C. MANUMBAS</b> — Treasurer, Capacuan Small-Scale Miners Association — mobile 0930-407-1626; and",
 "<b>MR. RAMON R. RAMOS</b> — Secretary, Capacuan Small-Scale Miners Association — mobile ____________________, email monraso1959@gmail.com,",
 ("— each of them to act singly, as my authorized representatives to <b>facilitate, file, follow up and receive "
  "documents, clearances, permits and licenses</b> on my behalf, for:"),
 ("<b>1. AVI GOLD PROCESSING PLANT</b>, my sole proprietorship registered with the Department of Trade and "
  "Industry under Business Name No. 8480852, with plant site at Lot 4, Psu-143364, OCT No. P-1616, Barangay "
  "Santa Rosa Sur, Jose Panganiban, Camarines Norte; and"),
 ("<b>2. The CAPACUAN SMALL-SCALE MINERS ASSOCIATION</b>, Purok 4, Barangay Capacuan, Paracale, Camarines "
  "Norte (DOLE Registration No. R0500-CN-0417-019-17), in connection with its Minahang Bayan petition and its "
  "small-scale mining permit applications."),
 ("For these purposes they may secure and submit application forms and letters, sign the receiving copies and "
  "logbooks as the persons filing, pay the official fees against an order of payment and receive the official "
  "receipts, request and receive certified true copies and certifications, follow up any application, receive "
  "notices and the permits or certificates issued, and accompany your inspectors during ocular inspection."),
 ("This authority is limited to documents and permits. My representatives are <b>not</b> authorized to sign any "
  "contract or agreement, to sell, mortgage or encumber any property or right of mine, to receive or hold money "
  "for my account other than official receipts, or to sign any affidavit or sworn statement in my place — those "
  "I sign personally."),
 ("This authorization is valid until <b>31 December 2027</b> unless I revoke it earlier in writing. Photocopies "
  "of my valid identification card and of my representatives' identification cards are attached."),
 "Very truly yours,", "", "",
 "<b>ALLAN VILLAFRIA INOCALLA</b>",
 "Proprietor, AVI Gold Processing Plant · TIN 200-031-253",
 "Mobile 0917 155 4782 · shiraction2@gmail.com",
 "ID presented: Driver’s Licence No. X01-11-002246 (LTO), valid to 12 January 2033", "", "",
 "<b>Specimen signatures of the authorized representatives:</b>", "",
 "______________________________                              ______________________________",
 "ROMEO C. MANUMBAS                                 RAMON R. RAMOS",
 "ID No. ______________________                              ID No. ______________________",
]

def clean(t): return t.replace('<b>','').replace('</b>','')

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.9); s.left_margin = s.right_margin = Inches(1.0)
doc.styles['Normal'].font.name = 'Times New Roman'; doc.styles['Normal'].font.size = Pt(11)
def para(t, bold=False, align=None, size=None, space=6):
    p = doc.add_paragraph(); r = p.add_run(clean(t)); r.bold = bold or t.startswith('<b>')
    if size: r.font.size = Pt(size)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'j': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space); return p
for k, t in HEAD: para(t, bold=(k == 'b'), align='c', size=(13 if k == 'b' else 10), space=0)
para("")
for t in BODY: para(t, align='j' if len(clean(t)) > 70 else None, space=4)
doc.save("AUTHORIZATION_LETTER_MANUMBAS_RAMOS_2026-09.docx")

ss = getSampleStyleSheet()
J = ParagraphStyle('J', parent=ss['Normal'], fontName='Times-Roman', fontSize=10.5, leading=13, alignment=TA_JUSTIFY, spaceAfter=4)
L = ParagraphStyle('L', parent=J, alignment=0, spaceAfter=2)
C = ParagraphStyle('C', parent=J, alignment=TA_CENTER, spaceAfter=0, fontSize=10)
CB = ParagraphStyle('CB', parent=C, fontName='Times-Bold', fontSize=13)
story = [Paragraph(t, CB if k == 'b' else C) for k, t in HEAD]
story.append(Spacer(1, 10))
for t in BODY:
    story.append(Paragraph(t or "&nbsp;", J if len(clean(t)) > 70 else L))
SimpleDocTemplate("AUTHORIZATION_LETTER_MANUMBAS_RAMOS_2026-09.pdf", pagesize=A4, leftMargin=2.3*cm,
                  rightMargin=2.3*cm, topMargin=1.5*cm, bottomMargin=1.3*cm).build(story)
print("built letter")
