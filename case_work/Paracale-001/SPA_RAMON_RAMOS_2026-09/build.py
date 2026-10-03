# -*- coding: utf-8 -*-
"""Standalone SPA — Allan V. Inocalla → Ramon R. Ramos. Run: python3 build.py
One instrument, ready to print and notarise. Paracale-001. Nothing filed."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER

BN   = "AVI GOLD PROCESSING PLANT"
DTI  = ("Business Name No. 8480852, Regional scope (Region V), valid 11 September 2026 to 11 September 2031, "
        "registered with the Department of Trade and Industry in my name")
SITE = ("Lot 4, Psu-143364, covered by Original Certificate of Title No. P-1616 of the Registry of Deeds for "
        "Camarines Norte, Barangay Santa Rosa Sur, Jose Panganiban, Camarines Norte")
LIC  = "Driver’s Licence No. X01-11-002246, Land Transportation Office, valid until 12 January 2033"

BODY = [
 "REPUBLIC OF THE PHILIPPINES )",
 "_________________________________ ) S.S.", "",
 "<c><b>SPECIAL POWER OF ATTORNEY</b></c>", "",
 "KNOW ALL MEN BY THESE PRESENTS:", "",
 ("I, <b>ALLAN VILLAFRIA INOCALLA</b>, Filipino, of legal age, with residence at Purok 4, Barangay Capacuan, "
  "Paracale, Camarines Norte, holder of Taxpayer Identification Number 200-031-253, and sole proprietor of "
  "<b>" + BN + "</b>, a sole proprietorship registered under " + DTI + ", with plant site at " + SITE +
  ", do hereby NAME, CONSTITUTE and APPOINT —"),
 ("<b>RAMON R. RAMOS</b>, Filipino, of legal age, Secretary of the Capacuan Small-Scale Miners Association, "
  "with residence at Purok 4, Barangay Capacuan, Paracale, Camarines Norte, mobile ____________________, "
  "email monraso1959@gmail.com,"),
 ("— as my true and lawful <b>ATTORNEY-IN-FACT</b>, for me and in my name, place and stead, <b>for the sole "
  "and limited purpose of facilitating the documents, clearances, permits and licences of " + BN + " and of "
  "its plant site</b>, with full power and authority to do and perform the following acts:"),
 ("<b>1. Obtain and lodge papers.</b> To secure application forms, checklists, requirements lists, orders of "
  "payment and assessment forms; to file, submit and lodge applications, letters, requests, petitions, "
  "compliance submissions, manifestations and follow-up letters in my name; and to sign the receiving copies, "
  "receiving logbooks and transmittal records <b>as the person filing on my behalf</b>."),
 ("<b>2. Pay official fees.</b> To pay filing, processing, registration, clearance, inspection and "
  "certification fees and other official charges against an official order of payment, and to demand, receive "
  "and keep the official receipts, which shall be issued in my name or in the name of " + BN + "."),
 ("<b>3. Obtain records.</b> To request and receive certified true copies, certifications, certified plans and "
  "maps, resolutions, endorsements, clearances, notices, orders and other documents relating to " + BN + ", to "
  "the plant site, and to my applications."),
 ("<b>4. Follow up and cure deficiencies.</b> To follow up the status of any application; to receive notices, "
  "orders, assessments and returns; to be informed of and to report to me any deficiency, defect or additional "
  "requirement; and to submit the curing documents that I furnish."),
 ("<b>5. Receive the issued instruments.</b> To receive the permits, licences, certificates, clearances and "
  "approvals issued, and to deliver the same to me intact."),
 ("<b>6. Attend and accompany.</b> To attend conferences, briefings, pre-assessment meetings and public "
  "consultations on my behalf; to accompany inspectors and evaluators to the plant site and to give them "
  "access to it; and to arrange the logistics of any ocular inspection."),
 ("<b>7. Before which offices.</b> The foregoing authority may be exercised before, among others: the "
  "<b>Mines and Geosciences Bureau, Regional Office No. V</b> (Rawis, Legazpi City) and the MGB Central "
  "Office; the <b>Environmental Management Bureau, Regional Office No. V</b>; the <b>Department of Environment "
  "and Natural Resources</b>, Regional Office V, PENRO and CENRO Camarines Norte; the <b>Provincial Mining "
  "Regulatory Board of Camarines Norte</b>; the <b>Provincial Government of Camarines Norte</b> and the "
  "Sangguniang Panlalawigan; the <b>Municipality of Jose Panganiban</b> (Office of the Mayor, Sangguniang "
  "Bayan, Business Permits and Licensing Office, Assessor, Treasurer, Engineering, Zoning, MENRO, "
  "Health/Sanitation) and the <b>Bureau of Fire Protection</b>; <b>Barangay Santa Rosa Sur</b>, Jose "
  "Panganiban, and <b>Barangay Capacuan</b>, Paracale; the <b>Department of Trade and Industry</b>; the "
  "<b>Bureau of Internal Revenue</b> in the Revenue District Office having jurisdiction; the <b>Registry of "
  "Deeds for Camarines Norte</b> and the <b>Land Registration Authority</b>, for certified copies of title "
  "only; the <b>National Mapping and Resource Information Authority</b>; the <b>National Commission on "
  "Indigenous Peoples</b>; and any other government office, agency, bureau or local government unit whose "
  "action is required for the same purpose."),
 ("<b>8. Specific transactions before the Bureau of Internal Revenue.</b> Without limiting the generality of "
  "the foregoing, and for the purpose of the Revenue requirement that specific transactions be named, my "
  "attorney-in-fact is expressly authorised: to file <b>BIR Form No. 1901</b> (Application for Registration) "
  "for " + BN + " together with its Part X payment order on <b>BIR Form No. 0605</b>; to file <b>BIR Form No. "
  "1905</b> for the transfer of my registration records or the update of my registration details, should the "
  "District Office require it; to pay the documentary stamp tax and the cost of invoices and to receive the "
  "official receipts; to purchase <b>BIR Printed Invoices</b> at the New Business Registrant Counter; and to "
  "receive my <b>Certificate of Registration (BIR Form No. 2303)</b>, the invoice booklets and the received "
  "copies of the foregoing, and deliver them to me."),
 ("<b>EXPRESS LIMITS ON THIS AUTHORITY.</b> This Special Power of Attorney is strictly limited to the "
  "facilitation of documents and permits described above. My attorney-in-fact is <b>NOT</b> authorised, and "
  "shall have no power whatsoever:"),
 "(a) to sell, convey, transfer, assign, mortgage, pledge, lease, or in any manner encumber or dispose of any land, title, permit, licence, equipment or other asset of mine or of " + BN + ";",
 "(b) to enter into, sign or agree to any contract, memorandum of agreement, joint venture, ore supply or offtake agreement, loan, investment, partnership or financial commitment, or to bind me in any obligation;",
 "(c) to receive, collect, hold or disburse any money or property for my account, except official receipts and the documents named above; and not to borrow, advance or lend money in my name;",
 "(d) to execute, sign or swear to any affidavit, sworn application, sworn statement, undertaking or declaration that requires my personal knowledge or my personal oath — those instruments I shall sign myself;",
 "(e) to waive, compromise, settle, quitclaim or release any right, claim or cause of action of mine;",
 "(f) to appoint a substitute or to delegate any part of this authority to any other person;",
 "(g) to represent me before any court, prosecutor’s office, or in any adversarial, criminal or quasi-judicial contest, or to perform any act that constitutes the practice of law; and",
 "(h) to give, offer or pay anything to any public officer other than the official fees supported by an official receipt.",
 ("<b>ACCOUNTING.</b> My attorney-in-fact shall turn over to me every document received and every official "
  "receipt within a reasonable time from receipt, and shall render an accounting of all sums advanced or paid."),
 ("<b>TERM AND REVOCATION.</b> This Special Power of Attorney shall take effect upon its notarisation and "
  "shall remain in force until <b>31 December 2027</b>, unless sooner revoked. It is revocable by me at any "
  "time and for any reason, by written notice to my attorney-in-fact, and I undertake to furnish a copy of the "
  "revocation to the offices before which this instrument has been presented. Any act done in good faith by a "
  "third person in reliance on this instrument before receipt of the revocation shall be respected."),
 ("<b>SCOPE NOTE.</b> This instrument carries my own authority as owner of " + BN + ". It does not, and is not "
  "intended to, govern the filings of the <b>Capacuan Small-Scale Miners Association</b>, in which Mr. Ramos "
  "holds the office of Secretary in his own right; those are made on the authority of the Association’s own "
  "resolution and his certification as its Secretary."),
 ("HEREBY GIVING AND GRANTING unto my said attorney-in-fact full power and authority to do and perform every "
  "act requisite and necessary to carry out the limited purposes above, as fully to all intents and purposes "
  "as I might or could lawfully do if personally present, <b>save only the acts expressly withheld above</b>, "
  "and hereby ratifying and confirming all that my said attorney-in-fact shall lawfully do or cause to be done "
  "under and by virtue hereof."),
 "", "",
 ("IN WITNESS WHEREOF, I have hereunto set my hand this ____ day of ______________ 2026 at "
  "______________________________, Philippines."),
 "", "", "",
 "<b>ALLAN VILLAFRIA INOCALLA</b>",
 "Principal · Proprietor, " + BN + " · TIN 200-031-253",
 LIC, "", "",
 "CONFORME — accepting the trust and the limits written above:", "", "", "",
 "<b>RAMON R. RAMOS</b>",
 "Attorney-in-Fact · ID: ______________________ No. ______________", "", "",
 "SIGNED IN THE PRESENCE OF:", "",
 "______________________________                    ______________________________", "",
 "REPUBLIC OF THE PHILIPPINES )",
 "_________________________________ ) S.S.", "",
 ("BEFORE ME, a Notary Public for and in _________________________, Philippines, this ____ day of "
  "______________ 2026, personally appeared <b>ALLAN VILLAFRIA INOCALLA</b>, who is personally known to me "
  "and who exhibited to me as competent evidence of identity his <b>" + LIC + "</b>, and who acknowledged to "
  "me that the foregoing instrument is his free and voluntary act and deed."),
 ("This instrument, consisting of ______ pages including this page on which this acknowledgment is written, "
  "has been signed by the principal and his instrumental witnesses on each and every page hereof."),
 "", "WITNESS MY HAND AND SEAL on the date and at the place first above written.", "", "", "",
 "NOTARY PUBLIC",
 "Doc. No. ______; Page No. ______; Book No. ______; Series of 2026.",
]

def clean(t): return t.replace('<b>','').replace('</b>','').replace('<c>','').replace('</c>','')

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.85); s.left_margin = s.right_margin = Inches(1.0)
doc.styles['Normal'].font.name = 'Times New Roman'; doc.styles['Normal'].font.size = Pt(11.5)
for t in BODY:
    p = doc.add_paragraph(); r = p.add_run(clean(t))
    r.bold = t.strip().startswith('<b>')
    if t.startswith('<c>'): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif len(clean(t)) > 70: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
doc.save("SPA_RAMON_RAMOS_2026-09.docx")

ss = getSampleStyleSheet()
J = ParagraphStyle('J', parent=ss['Normal'], fontName='Times-Roman', fontSize=11, leading=14.5, alignment=TA_JUSTIFY, spaceAfter=6)
L = ParagraphStyle('L', parent=J, alignment=0, spaceAfter=3)
C = ParagraphStyle('C', parent=J, alignment=TA_CENTER, spaceAfter=3, fontName='Times-Bold', fontSize=13)
story = []
for t in BODY:
    st = C if t.startswith('<c>') else (J if len(clean(t)) > 70 else L)
    story.append(Paragraph(t.replace('<c>','').replace('</c>','') or "&nbsp;", st))
SimpleDocTemplate("SPA_RAMON_RAMOS_2026-09.pdf", pagesize=A4, leftMargin=2.3*cm, rightMargin=2.3*cm,
                  topMargin=2*cm, bottomMargin=1.8*cm, title="Special Power of Attorney - Ramon R. Ramos").build(story)
print("built SPA_RAMON_RAMOS_2026-09.docx + .pdf")
