# Builds the barangay mediation filing set for Allan V. Inocalla v. Mauro C. Bombita Sr. et al.
#   01_COMPLAINT_AFFIDAVIT  (docx + pdf, 8.5x13)   -> signed by Allan, received by the barangay
#   02_ANNEXES_A-F.pdf                             -> annex covers + verified source pages
#   00_FILING_SET_bound.pdf                        -> affidavit + annexes, one PDF (print 3 sets)
#   03_ALLAN_POCKET_CARD (docx + pdf)              -> PRIVATE. Not filed. Not shown to anyone.
# Run: python3 build.py   (needs python-docx, PyMuPDF, soffice)
import os, subprocess, fitz
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "annex_src")
W, H = 612, 936  # 8.5 x 13 in, points

# ----------------------------------------------------------------------------- helpers
def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(13)
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.9)
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    return d

def para(d, text, bold=False, align=None, size=None, italic=False, indent=None):
    p = d.add_paragraph()
    # **bold** segments
    parts = text.split("**")
    for i, part in enumerate(parts):
        r = p.add_run(part); r.bold = bold or (i % 2 == 1); r.italic = italic
        if size: r.font.size = Pt(size)
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r": p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "l": p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent: p.paragraph_format.left_indent = Inches(indent)
    return p

def to_pdf(docx_path):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", HERE, docx_path],
                   check=True, capture_output=True)
    return docx_path[:-5] + ".pdf"

# ----------------------------------------------------------------------------- 01 affidavit
def build_affidavit():
    d = new_doc()
    para(d, "RECEIVED by Barangay ______________________\nDate/Time: ____________  By: ______________", align="r", size=9)
    para(d, "Republic of the Philippines\nProvince of Camarines Norte\nMunicipality of Jose Panganiban\nBARANGAY [______________________]\nOFFICE OF THE LUPONG TAGAPAMAYAPA",
         bold=True, align="c")
    para(d, "ALLAN VILLAFRIA INOCALLA, for himself and for the benefit of his co-heirs,\nComplainant,", align="l")
    para(d, "- versus -", align="c")
    para(d, "Barangay Case No. ________\nFor: Recovery of possession; stoppage and removal\nof unauthorized rod mill; accounting", align="r")
    para(d, "MAURO C. BOMBITA, SR., MAURO D. BOMBITA, JR., and all persons claiming rights under them,\nRespondents.", align="l")
    para(d, "x - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - - x", align="c")
    para(d, "COMPLAINT-AFFIDAVIT", bold=True, align="c", size=14)

    para(d, "I, **ALLAN VILLAFRIA INOCALLA**, [71] years old, Filipino, [separated], and residing at Purok 4, Barangay Capacuan, "
            "Paracale, Camarines Norte, after being duly sworn, state:")

    paras = [
        # Who I am / standing
        "I am a son and compulsory heir of the late spouses **Vicente Inocalla, Sr.** and **Beatriz Villafria Inocalla**, both deceased. "
        "I am also a brother of the late **Senen V. Inocalla**, who died on 12 February 2021. [My sister Senen never married and left no child.] "
        "Her properties therefore pass to her brothers and sisters and to the children of those who died before her, including me.",

        "I bring this complaint as one of the heirs and co-owners, for my own interest and for the benefit of all my co-heirs. "
        "Any co-owner may act to recover and protect common property for the benefit of all.",

        # The land
        "The land involved is:",
        "(a) **Lot No. 905, Pls-819-D**, a coconut land of about **7.8844 hectares** at Barangay San Rafael, Jose Panganiban, Camarines Norte, "
        "declared for taxation under **ARP No. 021-00470 (GR'23)** in the name of my late sister Senen V. Inocalla, as certified by the "
        "Office of the Provincial Assessor on 4 June 2024 (**Annex \"A\"**); and",
        "(b) the portion occupied by respondents of the land covered by **TCT No. T-2194** (13.3690 hectares), registered in the name of "
        "my father Vicente Inocalla married to Beatriz Villafria. In the judicial partition among us heirs, this property was "
        "assigned to Senen (**Annex \"B\"**). A certified copy of TCT No. T-2194 is available for inspection.",

        # How respondents came to be there
        "Respondents live at Purok 5, Barangay Sta. Rosa Sur, Jose Panganiban. They occupy about [five (5)] hectares of Lot 905 and "
        "about [2,000 square meters] of the land under TCT No. T-2194.",

        "Of my own knowledge, the Bombita family came to the land as **workers of my late brother Vicente Inocalla, Jr.**, who employed them in "
        "his coconut harvesting and small-scale mining. My parents never made them tenants. No leasehold contract was ever made with my parents "
        "or with us heirs. [Neither I nor, to my knowledge, any of my co-heirs has ever received any rental or share of the harvest from respondents.]",

        "As early as **11 February 2012**, my family met with barangay officials of Sta. Rosa Sur and Capacuan and a surveyor at the home of my sister "
        "Senen. The meeting took up, among other matters, \"a trespassing/eviction issue of Mauro Bombita and several of his family members\" "
        "(**Annex \"C\"**).",

        "In 1979, my father swore that the land covered by TCT No. T-2194 \"is not tenanted; nor portion thereof is devoted to rice or corn "
        "cultivation.\" That statement is annotated on the title. By Order dated **14 April 2008** in A-9999-05-EXE-096-03, the Secretary of "
        "Agrarian Reform **excluded the land under TCT No. T-2194 from agrarian reform coverage as mineral land** (**Annex \"D\"**).",

        # The prior case — disclosed, not conceded
        "I am aware that respondent Mauro C. Bombita, Sr. filed a case claiming to be a tenant against my late sister Senen alone "
        "(PARAD Case No. D-0502-RL-0015-2016 / DARAB Case Nos. 19415 and 19415-A, later CA-G.R. SP No. 167314). "
        "The estate of my parents and we, the other heirs, were not parties to that case. We do not recognize any tenancy by respondents binding "
        "on us, and we reserve all our rights. In any event, **nothing in that case allows respondents to build or operate an ore-processing mill, "
        "to mine, or to use the land for anything other than farming.**",

        # The rod mill — the present wrong
        "Since about [month/year], respondents, and persons acting with them, have been **constructing [and operating] a rod mill for grinding and "
        "processing gold ore** on Lot 905, at about [describe location / landmark / GPS]. They did this **without the knowledge or consent of the "
        "heirs**. [To my knowledge, they hold no permit for it from the Mines and Geosciences Bureau, the Environmental Management Bureau, or the "
        "Municipality.] Photographs taken on [date] are attached (**Annex \"E\"**).",

        "A rod mill is an ore-processing facility, not farming. Whatever respondents claim, the heirs never consented to any structure or "
        "mineral processing on the land. [Respondents have also [mined / dug tunnels / cut trees / brought in other persons to work or stay "
        "on the land] — describe with dates.]",

        "[Respondents have also prevented us heirs from harvesting coconuts on the land, claiming they own the right to it — give dates and "
        "incidents.]",

        # What I ask
        "I respectfully ask the Lupon, through mediation, to require respondents to:",
        "(1) **immediately STOP** constructing and operating the rod mill, and **REMOVE** the mill, its equipment, tanks and tailings from the land "
        "within fifteen (15) days;",
        "(2) **STOP** all mining, tunneling, tree-cutting and other non-farming use of the land;",
        "(3) **VACATE and peacefully SURRENDER** to the heirs the portions of Lot 905 and of the land under TCT No. T-2194 that they occupy, "
        "within a period to be agreed;",
        "(4) until they vacate, **not bring in** other persons, **not build** any new structure, and **not obstruct** the heirs' entry and harvest; and",
        "(5) **render a written accounting** of all coconuts, copra and other produce, and of any ore, taken from the land from [September 2023] to date.",
        "I also respectfully ask the Barangay to (a) conduct an **ocular inspection** of the rod mill and record what it finds, and (b) state "
        "whether any **barangay clearance or permit** was issued for its construction or operation.",

        "If no settlement is reached, I ask that a **Certification to File Action** be issued so that the heirs may bring the proper case "
        "before the proper court or agency.",

        "I execute this affidavit to attest to the truth of the foregoing, for the barangay proceedings, and for any other lawful purpose.",
    ]
    n = 0
    for t in paras:
        if t.startswith("(") :
            para(d, t, indent=0.5)
        elif t.endswith(":") and len(t) < 90:
            n += 1; para(d, f"{n}. {t}")
        else:
            n += 1; para(d, f"{n}. {t}")

    for p_ in d.paragraphs[-3:]: p_.paragraph_format.keep_with_next = True
    p_ = para(d, "IN WITNESS WHEREOF, I sign this on ______ September 2026 at ______________________, Camarines Norte.")
    p_.paragraph_format.keep_with_next = True
    para(d, "\n\n______________________________\nALLAN VILLAFRIA INOCALLA\nAffiant\n[ID type and No.: ____________________]", align="r")
    para(d, "SUBSCRIBED AND SWORN to before me this ______ day of September 2026 at ______________________, Camarines Norte, "
            "affiant exhibiting to me his [ID] No. ____________________.")
    para(d, "\n______________________________\n[Punong Barangay / Notary Public]", align="r")
    para(d, "Annexes: \"A\" Provincial Assessor Certification (4 Jun 2024) · \"B\" Judicial Partition, Civil Case No. 5825, RTC Br. 41 Daet · "
            "\"C\" Summation of barangay meeting (11 Feb 2012) · \"D\" DAR Secretary Order, A-9999-05-EXE-096-03 (14 Apr 2008) · "
            "\"E\" Photographs of rod mill · \"F\" Certificate of Death of Senen V. Inocalla", size=9)

    out = os.path.join(HERE, "01_COMPLAINT_AFFIDAVIT_Allan_Inocalla_v_Bombita.docx")
    d.save(out); return to_pdf(out)

# ----------------------------------------------------------------------------- 02 annexes
def cover(doc, letter, title, lines):
    p = doc.new_page(width=W, height=H)
    p.insert_textbox(fitz.Rect(60, 300, W - 60, 380), f'ANNEX "{letter}"', fontsize=34, fontname="tibo", align=1)
    p.insert_textbox(fitz.Rect(60, 400, W - 60, 470), title, fontsize=15, fontname="tibo", align=1)
    p.insert_textbox(fitz.Rect(80, 480, W - 80, 700), "\n".join(lines), fontsize=11, fontname="tiro", align=1)

def add_pdf_pages(doc, fn, pages):
    s = fitz.open(os.path.join(SRC, fn))
    for i in pages:
        p = doc.new_page(width=W, height=H)
        p.show_pdf_page(fitz.Rect(24, 24, W - 24, H - 24), s, i, keep_proportion=True)

def add_image(doc, fn):
    p = doc.new_page(width=W, height=H)
    p.insert_image(fitz.Rect(24, 24, W - 24, H - 24), filename=os.path.join(SRC, fn), keep_proportion=True)

def placeholder(doc, letter, title, lines):
    cover(doc, letter, title, lines)
    p = doc.new_page(width=W, height=H)
    p.draw_rect(fitz.Rect(60, 60, W - 60, H - 60), color=(0, 0, 0), width=1)
    p.insert_textbox(fitz.Rect(80, 420, W - 80, 520), f'[ATTACH ANNEX "{letter}" HERE]\n{title}', fontsize=14, fontname="tibo", align=1)

def build_annexes():
    doc = fitz.open()
    cover(doc, "A", "Certification, Office of the Provincial Assessor, Camarines Norte",
          ["4 June 2024 — Control No. 0429-0219",
           "ARP No. 021-00470 (GR'23), San Rafael, Jose Panganiban, 7.8844 has., Cocoland",
           "Declared owner: Senen V. Inocalla"])
    add_image(doc, "1493_assessor_cert.jpg")
    cover(doc, "B", "Judicial Partition — Civil Case No. 5825, RTC Branch 41, Daet",
          ["Francisco Inocalla, Vicente Inocalla Jr. and Cipriana Inocalla v. Jasper [Casper] Inocalla, et al.",
           "Parcel 12 (TCT No. T-2194, 13.3690 has.) awarded to Senen",
           "Copy as received by the RTC (6 pages)"])
    add_pdf_pages(doc, "671_partition.pdf", range(6))
    cover(doc, "C", "Summation of meeting, 11 February 2012",
          ["Inocalla family with Brgy. Councilor Arnel Competente (Sta. Rosa Sur),",
           "Brgy. Councilor Luis Salen (Capacuan) and surveyor Edgar S. Gerio",
           "Re: property line and the trespassing/eviction issue of Mauro Bombita"])
    add_image(doc, "1405_2012_meeting.jpg")
    cover(doc, "D", "Order of the Secretary of Agrarian Reform, 14 April 2008",
          ["A-9999-05-EXE-096-03 — Heirs of Vicente Inocalla, Sr., rep. by Casper Inocalla",
           "Excluding TCT No. T-2194 and other lots from CARP coverage as mineral land",
           "(6 pages)"])
    add_pdf_pages(doc, "1611_loan_with_dar_order.pdf", range(8, 14))
    placeholder(doc, "E", "Photographs of the rod mill on Lot 905",
                ["Each photo: date taken, place, and what it shows, initialed by affiant"])
    placeholder(doc, "F", "Certificate of Death of Senen V. Inocalla (PSA or LCR copy)",
                ["Attach if available; otherwise mark \"to be submitted\""])
    out = os.path.join(HERE, "02_ANNEXES_A-F.pdf"); doc.save(out, garbage=3, deflate=True); return out

# ----------------------------------------------------------------------------- 03 pocket card (private)
def build_card():
    d = new_doc()
    para(d, "ALLAN — MEDIATION GAME PLAN (PRIVATE — do not hand over, do not read aloud)", bold=True, align="c", size=13)
    sec = [
        ("THE POINT OF TOMORROW", [
            "The barangay cannot eject them. What it CAN do is make an **official record**. Leave with five things:",
            "1) **Their admissions in the minutes** (see QUESTIONS).  2) **An ocular inspection** of the mill by barangay officials.  "
            "3) If possible, a signed **freeze on the mill** (Interim Undertaking, exact text only).  4) **What they want in order to leave** — listen, do not answer.  "
            "5) **A copy of the minutes** and the next setting date (or a certification to file action).",
            "Tone: calm owner protecting family land. Not a vendetta. The Punong Barangay is also deciding your AVI endorsement — be the reasonable one in the room.",
        ]),
        ("QUESTIONS — ask the Punong Barangay to put these to them and WRITE THE ANSWERS in the minutes", [
            "Every answer helps us. A refusal to answer is also recorded.",
            "a) \"Who owns this land?\" (In his 2016 DARAB Answer Mauro Sr. admitted Senen/Inocalla ownership. If he says Inocalla → locked in. If he says \"us\" → he contradicts his own case.)",
            "b) \"Who built the rod mill, when, and who owns and finances it?\"",
            "c) \"Did any Inocalla heir give permission for the mill? In writing?\"",
            "d) \"Do you have a barangay clearance, mayor's permit, MGB or EMB permit for it?\"",
            "e) \"How many hectares exactly do you occupy, and where are the boundaries?\"",
            "f) \"Have you paid any rental or share since 2020? To whom? Receipts?\" (If \"to nobody\" or to a non-heir → proof of non-payment to the heirs.)",
            "g) \"Who else lives or works on the land?\" (names)",
            "h) \"Have you paid the land tax or applied for any title or patent on this land?\"",
        ]),
        ("OCULAR INSPECTION", [
            "Ask: \"May the Barangay send officials or tanods to see the mill and note it in the record?\" A barangay official's report is neutral proof "
            "and safer than you taking photos near their houses. Do NOT go with the inspection team yourself.",
        ]),
        ("THE ONLY THING YOU MAY SIGN", [
            "The **Interim Undertaking** (separate sheet), and only word-for-word: no further construction or operation of the mill while the matter is pending, "
            "without admission by anyone. If they want to change ANY word, or add money, sharing, or \"tenant\" → do not sign. Say you will consult your co-heirs.",
            "**Nothing else.** A barangay settlement becomes as binding as a court judgment 10 days after signing.",
        ]),
        ("NEVER", [
            "Never accept money, coconuts or copra. Never agree to split harvests. Never call them \"tenant\", \"kasama\" or \"tumatao\" — say \"occupants\" or \"workers of my late brother Vicente Jr.\"",
            "Never offer an amount. If they name a price to leave, write it down, say \"I will bring it to my co-heirs,\" and stop.",
        ]),
        ("DO NOT REVEAL", [
            "That Lot 905 has only a tax declaration and no title. (If asked: \"It is declared in Senen's name — Annex A.\")",
            "The Supreme Court case status (we do not know it). Radj Gymson. The 2014 levy on T-2194. Our DARAB or MGB plans.",
            "Nothing about the murders, the NBI, guns or suspicions. Not one word.",
            "Your AVI plant. If they raise it: \"That is on my own titled land, OCT P-1616, and I am applying for the permits. This is about a mill built on the heirs' land without consent.\"",
        ]),
        ("SAY", [
            "\"The land is ours as heirs. Nobody gave consent to a mill. It must stop and be removed.\"",
            "\"The Bombitas came as workers of my brother Vicente Jr., not as tenants of my parents.\"",
            "\"The heirs were never parties to the case against Senen. We do not recognize it.\"",
        ]),
        ("IF THEY SAY...", [
            "\"We won in DARAB / Court of Appeals.\" → \"That case was only against Senen. The heirs were not parties. No ruling lets you build a mill.\"",
            "\"The barangay has no jurisdiction / it is agrarian / you live in Paracale.\" → Do not argue. \"Please note that in the minutes.\" (If the barangay has no jurisdiction, we can go straight to the proper office without a certification anyway.)",
            "They show a decision → ask for a copy or to photograph it. No comment.",
            "\"You also have a levy on T-2194.\" → \"Separate matter. It gives you no right to the land.\"",
            "They do not appear → ask that their absence be recorded and for the certification to file action.",
        ]),
        ("SAFETY", [
            "Go with at least one companion; sit near the door; leave together; do not talk outside. No photos of them. "
            "If threatened, ask that it be written in the minutes and report to PNP provincial, not the JP station.",
        ]),
        ("BRING / AFTER", [
            "3 signed sets (barangay, respondents, you with RECEIVED stamp). ID. 2 copies of the Interim Undertaking. A notebook. Certified T-2194 copy stays in your folder unless the Lupon asks.",
            "AFTER, same day, send Jonathan: received copy · minutes (or photo) · their answers to a–h · any price they named · next setting date · any threats.",
        ]),
    ]
    for head, items in sec:
        para(d, head, bold=True, size=12)
        for it in items: para(d, "• " + it, size=11, indent=0.2)
    out = os.path.join(HERE, "03_ALLAN_POCKET_CARD_PRIVATE.docx"); d.save(out); return to_pdf(out)

def build_undertaking():
    d = new_doc()
    para(d, "Republic of the Philippines\nMunicipality of Jose Panganiban, Camarines Norte\nBARANGAY [______________________]\nOFFICE OF THE LUPONG TAGAPAMAYAPA", bold=True, align="c")
    para(d, "Allan Villafria Inocalla, for himself and his co-heirs v. Mauro C. Bombita, Sr., Mauro D. Bombita, Jr., et al.      Barangay Case No. ________", align="l")
    para(d, "INTERIM UNDERTAKING (STATUS QUO)", bold=True, align="c", size=14)
    items = [
        "While this matter is pending before the Barangay and until it is finally resolved, respondents Mauro C. Bombita, Sr. and Mauro D. Bombita, Jr., "
        "for themselves and all persons acting with them, undertake **not to continue constructing, not to operate, and not to expand the rod mill** "
        "and related facilities located at [location], and not to build any new structure on the land subject of this complaint.",
        "This undertaking is **without admission** by any party of any claim of ownership, possession or tenancy. It does **not settle** the complaint, "
        "which remains pending. All parties reserve all their rights, claims and defenses before any court or agency.",
        "The parties shall keep the peace and shall not harass or threaten one another.",
    ]
    for i, t in enumerate(items, 1): para(d, f"{i}. {t}")
    para(d, "Signed on ______ September 2026 at Barangay ______________________, Jose Panganiban, Camarines Norte.")
    para(d, "\n______________________________          ______________________________\nMAURO C. BOMBITA, SR.                              MAURO D. BOMBITA, JR.\n\n"
            "______________________________\nALLAN VILLAFRIA INOCALLA", align="c")
    para(d, "Attested:\n\n______________________________\nPunong Barangay / Lupon Chairman", align="c")
    out = os.path.join(HERE, "04_INTERIM_UNDERTAKING_optional.docx"); d.save(out); return to_pdf(out)

if __name__ == "__main__":
    aff = build_affidavit()
    anx = build_annexes()
    bound = fitz.open()
    for f in (aff, anx): bound.insert_pdf(fitz.open(f))
    bound.save(os.path.join(HERE, "00_FILING_SET_bound.pdf"), garbage=3, deflate=True)
    card = build_card()
    und = build_undertaking()
    for f in (aff, anx, os.path.join(HERE, "00_FILING_SET_bound.pdf"), card, und):
        print(f, len(fitz.open(f)), "pages")
