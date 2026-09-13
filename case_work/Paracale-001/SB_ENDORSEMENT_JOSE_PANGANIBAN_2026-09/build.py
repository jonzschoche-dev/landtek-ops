# Builds the SB Jose Panganiban endorsement request letter + draft SB Resolution as DOCX and PDF. Run: python3 build.py
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER

DATE = "____ September 2026"
BN = "AVI GOLD PROCESSING PLANT"
BN_NO = "DTI Business Name No. 8480852 (Regional – Region V), valid 11 September 2026 to 11 September 2031"

head = [("b", BN), ("n", "Allan V. Inocalla, Proprietor"),
        ("n", "Purok 4, Barangay Capacuan, Paracale, Camarines Norte 4605 (residence) · Plant site: Lot 4, Psu-143364, OCT No. P-1616, Barangay Santa Rosa Sur, Jose Panganiban, Camarines Norte 4606"),
        ("n", "Mobile 0917 155 4782  ·  shiraction2@gmail.com")]
addr = [DATE, "", "HON. CASIMERO B. PADILLA JR.", "Municipal Vice Mayor and Presiding Officer",
        "and the Honorable Members of the Sangguniang Bayan", "Municipality of Jose Panganiban, Camarines Norte", "",
        "Thru: The Secretary to the Sangguniang Bayan", "",
        "Copy: HON. ARIEL M. NON, Municipal Mayor"]
subject = ("SUBJECT: REQUEST FOR A SANGGUNIANG BAYAN RESOLUTION ENDORSING THE APPLICATIONS OF AVI GOLD PROCESSING PLANT "
           "(ALLAN V. INOCALLA, PROPRIETOR) FOR A MINERAL PROCESSING PERMIT / MINERAL PROCESSOR'S LICENSE WITH DENR–MGB REGION V "
           "AND THE PROVINCIAL MINING REGULATORY BOARD, AND FOR AN ENVIRONMENTAL COMPLIANCE CERTIFICATE WITH EMB REGION V, FOR A "
           "MERCURY-FREE GOLD PROCESSING PLANT AT LOT 4, PSU-143364 (OCT NO. P-1616), BARANGAY SANTA ROSA SUR; AND RECOMMENDING THE "
           "DESIGNATION OF THE PLANT SITE AS A MINERAL PROCESSING ZONE")
body = [
    "Honorable Vice Mayor Padilla and Honorable Members of the Sanggunian:",
    (f"I respectfully request the Sangguniang Bayan of Jose Panganiban to adopt a Resolution expressing conformance with, and "
     f"endorsing, the applications of <b>{BN}</b> for (1) a Mineral Processing Permit with the DENR – Mines and Geosciences Bureau, "
     "Regional Office No. V, and/or a Custom Mill / Mineral Processor's License with the Provincial Mining Regulatory Board of "
     "Camarines Norte, and (2) an Environmental Compliance Certificate with the Environmental Management Bureau, Region V, for a "
     "mercury-free gold processing plant to be established on my own titled land in Barangay Santa Rosa Sur, this Municipality."),
    ("<b>The site.</b> Lot 4, Psu-143364, Original Certificate of Title No. P-1616 of the Registry of Deeds for Camarines Norte "
     "(Free Patent No. 225537 issued 10 May 1963), 15.2069 hectares, situated in Barangay Santa Rosa Sur and registered solely in "
     "my name, Allan Villafria Inocalla. The plant will occupy only a portion of the lot; the remainder stays in its present use."),
    (f"<b>The business.</b> {BN} is registered with the Department of Trade and Industry under {BN_NO}, issued to me. It will be a "
     "gravity-concentration plant of about ten (10) tonnes of ore per day, below ten thousand (10,000) tonnes per year: crushing "
     "and milling, classification, a centrifugal rougher and locally fabricated concentrating (shaking) tables, with the "
     "concentrate smelted to doré. <b>No mercury and no cyanide will be used at any stage.</b> All gold produced will be sold to "
     "the Bangko Sentral ng Pilipinas or its accredited traders. The plant will process ore from registered small-scale miners "
     "and lawful mining-rights holders of Jose Panganiban and the Paracale–Jose Panganiban gold district under written supply "
     "agreements, including the Minahang Bayan already declared in Barangay Santa Rosa Norte and the Minahang Bayan applied for "
     "by the Capacuan Small-Scale Miners Association over the Inocalla properties in Paracale–Jose Panganiban."),
    ("<b>Why the Sanggunian's resolution is needed.</b> The Mines and Geosciences Bureau requires, for a Mineral Processing Permit, "
     "\"project approval/endorsement by at least majority of the Sanggunian concerned\" (MGB Citizen's Charter 2026, Mineral "
     "Processing Permit, Other Requirements). Under Sections 26 and 27 of the Local Government Code, projects of this kind require "
     "prior consultation with, and the approval of, the Sanggunian where the project is located. Under DENR Administrative Order "
     "No. 2022-03, Section 15, Mineral Processing Zones for small-scale mineral processing are designated by the local government "
     "unit upon recommendation of the Provincial Mining Regulatory Board. This request follows the same course the Sangguniang "
     "Bayan of Paracale took in Resolution No. 140-2012 for our family's earlier processing plant at Barangay Capacuan."),
    "I therefore respectfully ask the Sangguniang Bayan to resolve:",
    (f"1. To express conformance with and endorse the application of Allan V. Inocalla / {BN} for a Mineral Processing Permit "
     "with DENR–MGB Region V and/or a Custom Mill / Mineral Processor's License with the Provincial Mining Regulatory Board of "
     "Camarines Norte, for the plant at Lot 4, Psu-143364 (OCT No. P-1616), Barangay Santa Rosa Sur, Jose Panganiban;"),
    ("2. To endorse the plant's application for an Environmental Compliance Certificate with EMB Region V and to certify that the "
     "project is compatible with the Municipality's land-use plan, or to direct the Municipal Planning and Development Office to "
     "issue the corresponding locational/zoning clearance upon compliance with its requirements;"),
    ("3. To recommend to the Provincial Mining Regulatory Board the designation of the plant site as a Mineral Processing Zone under "
     "Section 15 of DAO No. 2022-03, and to signify the Municipality's readiness to designate it as such upon the Board's "
     "recommendation; and"),
    ("4. To furnish certified copies of the Resolution to the undersigned, the Office of the Municipal Mayor, the Provincial Mining "
     "Regulatory Board / MGB Region V (Rawis, Legazpi City), EMB Region V, and the Sangguniang Barangay of Santa Rosa Sur."),
    ("<b>Commitments to the Municipality.</b> The plant will (a) use gravity concentration only, without mercury or cyanide; "
     "(b) contain all tailings in a lined tailings facility with no discharge to any watercourse, and comply with the conditions "
     "of the ECC and of the Board's environmental, community-development and safety programs; (c) give preference in employment to "
     "residents of Barangay Santa Rosa Sur and of Jose Panganiban; (d) register with the Municipality, pay the business taxes and "
     "fees due under the Municipal Revenue Code, and secure the Mayor's Permit before operating; and (e) submit to inspection by "
     "the Municipal Environment and Natural Resources Office and the Barangay at any time."),
    ("Enclosed are the Resolution of the Sangguniang Barangay of Santa Rosa Sur endorsing the project, a draft Resolution for the "
     "Sanggunian's consideration, and the supporting documents listed below. I am available to appear before the Committee on "
     "Environment and Natural Resources or at a public consultation at the Sanggunian's convenience."),
    "Respectfully yours,", "", "", "", "<b>ALLAN V. INOCALLA</b>", f"Proprietor, {BN}", "TIN 200-011-253", "",
    ("Enclosures: (1) Draft Sangguniang Bayan Resolution; (2) Sangguniang Barangay Santa Rosa Sur Resolution No. ____ s-2026 "
     "endorsing the project and Barangay Business Clearance; (3) DTI Certificate of Business Name Registration No. 8480852 and "
     "official receipt; (4) Certified true copy of OCT No. P-1616 and Tax Declaration No. ______; (5) Sketch plan / site "
     "development plan of the plant site showing the portion of Lot 4 to be used; (6) Project description (process flow, "
     "capacity, tailings management, water and power); (7) Two valid IDs of the applicant."),
    "Received by: ______________________  Position: ______________  Date/Time: ______________",
]

res_head = ["Republic of the Philippines", "Province of Camarines Norte", "Municipality of Jose Panganiban",
            "OFFICE OF THE SANGGUNIANG BAYAN", "",
            "EXCERPTS FROM THE MINUTES OF THE ____ REGULAR SESSION OF THE SANGGUNIANG BAYAN OF JOSE PANGANIBAN, CAMARINES NORTE, "
            "HELD ON ____________ 2026 AT THE MUNICIPAL SESSION HALL.", "",
            "PRESENT:  Hon. Casimero B. Padilla Jr. – Municipal Vice Mayor, Presiding Officer;  Hon. Don Isaiah M. Lim;  "
            "Hon. Maria Corazon G. Arenal;  Hon. Francesca Angela S. Napa;  Hon. Catherine D. Baylon;  Hon. Christian Anthony R. Abaño;  "
            "Hon. Ariel E. Villaflores;  Hon. Jan Michael O. Guzman;  Hon. Jason J. Arriola;  Hon. ____________________ – ABC President;  "
            "Hon. ____________________ – SK Federation President.", "ABSENT: ______________________", ""]
res_title = ("RESOLUTION NO. ______-2026 — RESOLUTION EXPRESSING CONFORMANCE WITH AND ENDORSING THE APPLICATIONS OF AVI GOLD PROCESSING "
             "PLANT (ALLAN V. INOCALLA, PROPRIETOR) FOR A MINERAL PROCESSING PERMIT / MINERAL PROCESSOR'S LICENSE WITH THE DENR – MINES AND "
             "GEOSCIENCES BUREAU REGION V AND THE PROVINCIAL MINING REGULATORY BOARD OF CAMARINES NORTE, AND FOR AN ENVIRONMENTAL COMPLIANCE "
             "CERTIFICATE WITH THE ENVIRONMENTAL MANAGEMENT BUREAU REGION V, FOR A MERCURY-FREE GOLD PROCESSING PLANT AT LOT 4, PSU-143364 "
             "(OCT NO. P-1616), BARANGAY SANTA ROSA SUR, JOSE PANGANIBAN, CAMARINES NORTE, AND RECOMMENDING THE DESIGNATION OF THE PLANT SITE "
             "AS A MINERAL PROCESSING ZONE")
res_body = [
    "Author: Hon. ______________________",
    ("WHEREAS, Republic Act No. 7076, the People's Small-Scale Mining Act of 1991, declares it the policy of the State to promote, "
     "develop, protect and rationalize viable small-scale mining activities in order to generate more employment opportunities and "
     "provide an equitable sharing of the nation's wealth and natural resources;"),
    ("WHEREAS, DENR Administrative Order No. 2022-03, the Revised Implementing Rules and Regulations of RA No. 7076, provides in "
     "Sections 14 and 15 that small-scale mineral processing shall be undertaken in custom mills within Mineral Processing Zones "
     "designated by the local government unit concerned upon recommendation of the Provincial Mining Regulatory Board, and that no "
     "mercury shall be used in mineral processing; and Republic Act No. 7942 and DENR Administrative Order No. 2010-21 provide for "
     "the issuance of Mineral Processing Permits by the Mines and Geosciences Bureau;"),
    ("WHEREAS, Sections 26 and 27 of Republic Act No. 7160, the Local Government Code of 1991, require prior consultation with and the "
     "approval of the Sanggunian concerned for projects that may affect the environment of the locality, and the Mines and "
     "Geosciences Bureau requires the endorsement of the majority of the Sanggunian concerned for applications for Mineral "
     "Processing Permits;"),
    ("WHEREAS, presented to this Body is the application of Mr. Allan V. Inocalla, proprietor of AVI GOLD PROCESSING PLANT (DTI "
     "Business Name No. 8480852), for a mercury-free, cyanide-free gold gravity-concentration plant of about ten (10) tonnes per day "
     "on a portion of Lot 4, Psu-143364, covered by Original Certificate of Title No. P-1616 of the Registry of Deeds for Camarines "
     "Norte, registered solely in his name, situated at Barangay Santa Rosa Sur, this Municipality, together with the endorsement of "
     "the Sangguniang Barangay of Santa Rosa Sur under its Resolution No. ____ s-2026, the sketch plan of the site, the project "
     "description and other supporting documents;"),
    ("WHEREAS, the plant will process ore from registered small-scale miners and lawful mining-rights holders of this Municipality "
     "and the Paracale–Jose Panganiban gold district, will sell all gold produced to the Bangko Sentral ng Pilipinas or its "
     "accredited traders, will give preference in employment to residents of Barangay Santa Rosa Sur and of this Municipality, and "
     "will contain all tailings within a lined tailings facility in compliance with the conditions of its Environmental Compliance "
     "Certificate;"),
    ("WHEREAS, a legitimate, mercury-free custom mill within the Municipality will provide the small-scale miners of Jose Panganiban "
     "a lawful and safe alternative to mercury-based processing, generate employment and local revenue, and support the "
     "Municipality's responsible small-scale mining program;"),
    ("NOW, THEREFORE, on motion of Hon. ______________________, duly seconded by Hon. ______________________, be it —"),
    ("RESOLVED, as it is hereby resolved, to express conformance with and endorse the application of Mr. Allan V. Inocalla / AVI GOLD "
     "PROCESSING PLANT for a Mineral Processing Permit with the DENR – Mines and Geosciences Bureau, Regional Office No. V, and/or a "
     "Custom Mill / Mineral Processor's License with the Provincial Mining Regulatory Board of Camarines Norte, for the gold "
     "processing plant at Lot 4, Psu-143364 (OCT No. P-1616), Barangay Santa Rosa Sur, Jose Panganiban, Camarines Norte, subject to "
     "compliance with mining laws, environmental laws, local ordinances and other existing applicable laws;"),
    ("RESOLVED FURTHER, to endorse the said plant's application for an Environmental Compliance Certificate with the Environmental "
     "Management Bureau, Region V, and to direct the Municipal Planning and Development Office to issue the corresponding "
     "locational/zoning clearance upon the applicant's compliance with its requirements;"),
    ("RESOLVED FURTHER, to recommend to the Provincial Mining Regulatory Board of Camarines Norte the designation of the said plant "
     "site as a Mineral Processing Zone pursuant to Section 15 of DENR Administrative Order No. 2022-03, and to signify the readiness "
     "of this Municipality to designate the same upon the Board's recommendation;"),
    ("RESOLVED FINALLY, to furnish copies of this Resolution to Mr. Allan V. Inocalla, the Office of the Municipal Mayor, the "
     "Provincial Mining Regulatory Board / DENR-MGB Region V, Rawis, Legazpi City, the Environmental Management Bureau Region V, the "
     "Municipal Planning and Development Office, the Municipal Environment and Natural Resources Office, and the Sangguniang Barangay "
     "of Santa Rosa Sur, for their information and appropriate action."),
    "UNANIMOUSLY APPROVED.", "",
    "I hereby certify to the correctness of the foregoing resolution.", "", "",
    "______________________________", "Secretary to the Sangguniang Bayan", "", "ATTESTED:", "", "",
    "HON. CASIMERO B. PADILLA JR.", "Municipal Vice Mayor / Presiding Officer", "", "APPROVED:", "", "",
    "HON. ARIEL M. NON", "Municipal Mayor", "",
    "(Draft prepared by the applicant for the Sanggunian's consideration, following the form of Sangguniang Bayan of Paracale "
     "Resolution No. 140-2012; members per DILG Region V Masterlist of Local Officials 2025–2028.)",
]

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.9); s.left_margin = s.right_margin = Inches(1.0)
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11.5)
def para(text, bold=False, align=None, size=None, space=6):
    p = doc.add_paragraph(); clean = text.replace('<b>','').replace('</b>','')
    r = p.add_run(clean); r.bold = bold or text.startswith('<b>')
    if size: r.font.size = Pt(size)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'j': p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space); return p
for k, t in head: para(t, bold=(k=='b'), align='c', size=(13 if k=='b' else 10.5), space=0)
para("")
for t in addr: para(t, bold=t.startswith("HON"), space=0)
para(""); para(subject, bold=True, align='j'); para("")
for t in body: para(t, align='j')
doc.add_page_break()
for t in res_head: para(t, bold=(t.isupper() and len(t) < 40), align=('c' if len(t) < 40 else 'j'), space=0)
para(res_title, bold=True, align='j'); para("")
for t in res_body: para(t, align='j' if len(t) > 60 else None, bold=(t.isupper() and 5 < len(t) < 40))
doc.save("SB_JOSE_PANGANIBAN_ENDORSEMENT_REQUEST_AVI_GOLD_PROCESSING_PLANT.docx")

ss = getSampleStyleSheet()
J = ParagraphStyle('J', parent=ss['Normal'], fontName='Times-Roman', fontSize=11, leading=14.5, alignment=TA_JUSTIFY, spaceAfter=6)
L = ParagraphStyle('L', parent=J, alignment=0, spaceAfter=0)
C = ParagraphStyle('C', parent=J, alignment=TA_CENTER, spaceAfter=0)
CB = ParagraphStyle('CB', parent=C, fontName='Times-Bold', fontSize=13)
B = ParagraphStyle('B', parent=J, fontName='Times-Bold')
story = []
for k, t in head: story.append(Paragraph(t, CB if k=='b' else C))
story.append(Spacer(1, 10))
for t in addr: story.append(Paragraph(("<b>%s</b>" % t) if t.startswith("HON") else (t or "&nbsp;"), L))
story.append(Spacer(1, 8)); story.append(Paragraph(subject, B)); story.append(Spacer(1, 4))
for t in body: story.append(Paragraph(t or "&nbsp;", J))
story.append(PageBreak())
for t in res_head: story.append(Paragraph(("<b>%s</b>" % t) if (t.isupper() and len(t) < 40) else (t or "&nbsp;"), C if len(t) < 40 else J))
story.append(Spacer(1, 6)); story.append(Paragraph(res_title, B)); story.append(Spacer(1, 4))
for t in res_body: story.append(Paragraph(("<b>%s</b>" % t) if (t.isupper() and 5 < len(t) < 40) else (t or "&nbsp;"), J))
SimpleDocTemplate("SB_JOSE_PANGANIBAN_ENDORSEMENT_REQUEST_AVI_GOLD_PROCESSING_PLANT.pdf", pagesize=A4,
                  leftMargin=2.3*cm, rightMargin=2.3*cm, topMargin=2*cm, bottomMargin=2*cm).build(story)
print("built")
